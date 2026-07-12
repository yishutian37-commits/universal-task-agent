from __future__ import annotations

from datetime import datetime
from pathlib import Path

from rag.errors import IngestError
from rag.loaders.base import LoadedDoc


def _metadata(source: str, file_type: str, text: str, **extra) -> dict:
    path = Path(source)
    return {
        "type": file_type,
        "title": path.stem,
        "size": len(text),
        "ingested_at": datetime.now().isoformat(timespec="seconds"),
        **extra,
    }


class PdfLoader:
    def load(self, source: str) -> LoadedDoc:
        try:
            from pypdf import PdfReader

            reader = PdfReader(source)
            pages = [str(page.extract_text() or "").strip() for page in reader.pages]
        except FileNotFoundError:
            raise
        except Exception as exc:
            raise IngestError(f"读取 PDF 失败 {source}: {exc}") from exc
        text = "\n\n".join(page for page in pages if page).strip()
        if not text:
            raise IngestError(f"PDF 没有可提取文本，可能是扫描件: {source}")
        return LoadedDoc(text=text, source=source, metadata=_metadata(source, "pdf", text, page_count=len(pages)))


class DocxLoader:
    def load(self, source: str) -> LoadedDoc:
        try:
            from docx import Document

            document = Document(source)
            blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append("\t".join(cells))
        except FileNotFoundError:
            raise
        except Exception as exc:
            raise IngestError(f"读取 DOCX 失败 {source}: {exc}") from exc
        text = "\n".join(blocks).strip()
        if not text:
            raise IngestError(f"DOCX 没有可提取文本: {source}")
        return LoadedDoc(text=text, source=source, metadata=_metadata(source, "docx", text))

