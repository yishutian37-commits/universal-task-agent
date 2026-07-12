from __future__ import annotations

from docx import Document

from rag.loaders.base import LoaderFactory
from rag.loaders.document_loaders import DocxLoader, PdfLoader


def test_document_loader_factory_supports_pdf_and_docx():
    factory = LoaderFactory.for_documents()
    assert isinstance(factory.get("notes.pdf"), PdfLoader)
    assert isinstance(factory.get("notes.docx"), DocxLoader)


def test_pdf_loader_extracts_pages(monkeypatch, tmp_path):
    source = tmp_path / "guide.pdf"
    source.write_bytes(b"fake")

    class Page:
        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

    class Reader:
        def __init__(self, path):
            assert str(path) == str(source)
            self.pages = [Page("第一页"), Page("第二页")]

    monkeypatch.setattr("pypdf.PdfReader", Reader)
    loaded = PdfLoader().load(str(source))

    assert loaded.text == "第一页\n\n第二页"
    assert loaded.metadata["type"] == "pdf"
    assert loaded.metadata["page_count"] == 2


def test_docx_loader_extracts_paragraphs_and_tables(tmp_path):
    source = tmp_path / "guide.docx"
    document = Document()
    document.add_paragraph("项目介绍")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    document.save(source)

    loaded = DocxLoader().load(str(source))

    assert "项目介绍" in loaded.text
    assert "字段\t说明" in loaded.text
    assert loaded.metadata["type"] == "docx"
